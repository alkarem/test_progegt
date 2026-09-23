"""إخراج التقارير: HTML للطباعة، وPDF (WeasyPrint بخط عربي مضمَّن)، وExcel (أرقام حقيقية)."""
import html
import io
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.reports.engine import Column, ReportData

FONTS = Path(__file__).parent / "fonts"
TITLE_ORG = "نظام مراقبة الاعتمادات والمصروفات الحكومية"


def _fmt(v, kind: str) -> str:
    if v is None or v == "":
        return ""
    if kind == "money":
        d = Decimal(v)
        s = f"{abs(d):,.3f}"
        return f"({s})" if d < 0 else s
    if kind == "percent":
        return f"{Decimal(v):.2f}"
    if kind == "date":
        return v.strftime("%d/%m/%Y") if isinstance(v, (date, datetime)) else str(v)
    return str(v)


def to_html(data: ReportData, *, generated_by: str, for_pdf: bool = False) -> str:
    esc = html.escape
    font_face = "".join(
        f"@font-face{{font-family:'Plex';src:url('{(FONTS / f).as_uri()}');font-weight:{w};}}"
        for f, w in (("IBMPlexSansArabic-Regular.ttf", 400), ("IBMPlexSansArabic-Bold.ttf", 700)))
    wide = len(data.columns) > 7
    head = "".join(f"<th>{esc(c.title)}</th>" for c in data.columns)
    body = []
    for r in data.rows:
        cells = []
        for c in data.columns:
            v = r.get(c.key)
            neg = c.kind == "money" and v is not None and Decimal(v) < 0
            cls = ("num" if c.kind in ("money", "percent", "int") else "") + (" neg" if neg else "")
            cells.append(f'<td class="{cls}">{esc(_fmt(v, c.kind))}</td>')
        body.append("<tr>" + "".join(cells) + "</tr>")
    totals = data.totals
    foot = ""
    if totals:
        tds = []
        for i, c in enumerate(data.columns):
            if c.key in totals:
                tds.append(f'<td class="num">{esc(_fmt(totals[c.key], c.kind))}</td>')
            else:
                tds.append("<td>الإجمالي</td>" if i == 0 else "<td></td>")
        foot = "<tfoot><tr>" + "".join(tds) + "</tr></tfoot>"
    params = "، ".join(f"{k}: {v}" for k, v in data.params.items() if v not in (None, "", []))
    now = datetime.now(UTC).astimezone()
    notes = "".join(f"<p class='note'>{esc(n)}</p>" for n in data.notes)
    return f"""<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>{esc(data.title)}</title>
<style>{font_face}
@page {{ size: A4 {'landscape' if wide else 'portrait'}; margin: 14mm 10mm 16mm 10mm;
  @bottom-center {{ content: "صفحة " counter(page) " من " counter(pages); font-family: Plex; font-size: 8pt; color:#555; }} }}
body {{ font-family: Plex, 'DejaVu Sans', sans-serif; font-size: 9pt; color: #111; margin: 0; }}
header {{ border-bottom: 2px solid #0F3D5E; margin-bottom: 8px; padding-bottom: 6px; }}
header .org {{ color: #0F3D5E; font-weight: 700; font-size: 10pt; }}
h1 {{ font-size: 14pt; margin: 4px 0; }}
.sub {{ color: #333; }}
table {{ width: 100%; border-collapse: collapse; }}
th {{ background: #0F3D5E; color: #fff; font-weight: 700; padding: 4px; border: 1px solid #0F3D5E; }}
td {{ padding: 3px 4px; border: 1px solid #c9ced6; }}
tr:nth-child(even) td {{ background: #f6f7f9; }}
tfoot td {{ font-weight: 700; background: #e8edf3 !important; border-top: 2px solid #0F3D5E; }}
td.num {{ text-align: left; direction: ltr; font-variant-numeric: tabular-nums; white-space: nowrap; }}
td.neg {{ color: #B42318; }}
thead {{ display: table-header-group; }} tfoot {{ display: table-row-group; }}
.note {{ color: #7a5a00; font-size: 8pt; }}
footer {{ margin-top: 8px; font-size: 7.5pt; color: #555; border-top: 1px solid #c9ced6; padding-top: 4px; }}
@media print {{ .noprint {{ display: none; }} }}
</style></head><body>
<header><div class="org">{TITLE_ORG}</div><h1>{esc(data.title)} <small>({esc(data.code)})</small></h1>
<div class="sub">{esc(data.subtitle or '')}</div></header>
{notes}
<table><thead><tr>{head}</tr></thead><tbody>{''.join(body) or f'<tr><td colspan="{len(data.columns)}">لا توجد بيانات.</td></tr>'}</tbody>{foot}</table>
<footer>أُنشئ في {now:%d/%m/%Y %H:%M} بواسطة {esc(generated_by)} — المعاملات: {esc(params) or '—'} —
عدد الأسطر: {len(data.rows)} — بصمة المحتوى SHA-256: <span dir="ltr">{data.fingerprint()[:32]}…</span></footer>
{'' if for_pdf else '<script class="noprint">window.onload=()=>window.print()</script>'}
</body></html>"""


def to_pdf(data: ReportData, *, generated_by: str) -> bytes:
    from weasyprint import HTML
    return HTML(string=to_html(data, generated_by=generated_by, for_pdf=True), base_url=str(FONTS)).write_pdf()


def to_xlsx(data: ReportData, *, generated_by: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = data.code
    ws.sheet_view.rightToLeft = True
    bold = Font(bold=True)
    ws.append([TITLE_ORG])
    ws.append([f"{data.title} ({data.code})"])
    ws.append([data.subtitle or ""])
    ws["A1"].font = Font(bold=True, color="0F3D5E")
    ws["A2"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append([c.title for c in data.columns])
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="0F3D5E")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for r in data.rows:
        ws.append([_cell(r.get(c.key), c) for c in data.columns])
    totals = data.totals
    if totals:
        ws.append([("الإجمالي" if i == 0 else None) if c.key not in totals else float(0) for i, c in enumerate(data.columns)])
        tr = ws.max_row
        first, last = header_row + 1, tr - 1
        for i, c in enumerate(data.columns, start=1):
            if c.key in totals:
                col = get_column_letter(i)
                ws.cell(tr, i).value = f"=SUM({col}{first}:{col}{last})" if last >= first else 0
            ws.cell(tr, i).font = bold
    for i, c in enumerate(data.columns, start=1):
        fmt = {"money": "#,##0.000;[Red](#,##0.000)", "percent": "0.00", "int": "0", "date": "dd/mm/yyyy"}.get(c.kind)
        for row in ws.iter_rows(min_row=header_row + 1, min_col=i, max_col=i):
            if fmt:
                row[0].number_format = fmt
        ws.column_dimensions[get_column_letter(i)].width = 16 if c.kind != "text" else 28
    meta = wb.create_sheet("المعاملات")
    meta.sheet_view.rightToLeft = True
    meta.append(["أُنشئ بواسطة", generated_by])
    meta.append(["وقت الإنشاء", datetime.now(UTC).astimezone().strftime("%d/%m/%Y %H:%M")])
    meta.append(["بصمة المحتوى SHA-256", data.fingerprint()])
    for k, v in data.params.items():
        meta.append([k, str(v)])
    for n in data.notes:
        meta.append(["ملاحظة", n])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v, c: Column):
    """أرقام حقيقية في Excel (وليس نصوصًا) — 09-reports §1."""
    if v is None or v == "":
        return None
    if c.kind in ("money", "percent"):
        return Decimal(v)
    if c.kind == "int":
        return int(v)
    if c.kind == "date" and isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v
