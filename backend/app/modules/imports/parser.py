"""قراءة سجل مراقبة الاعتماد (نموذج 23/6) وتصنيف صفوفه (10-excel-import §2–§3).

دوال نقية لا تكتب في قاعدة البيانات. الملف يُقرأ مرتين: مرة للمعادلات (لكشف المعادلات
المستبدلة بقيم ثابتة) ومرة للقيم المحسوبة (للمقارنة فقط؛ لا تدخل أي رصيد).
"""
import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.shared.arabic import normalize

RE_ITEM = re.compile(r"البند\s*:\s*(\d+)\s*/\s*(\d+)")
RE_YEAR = re.compile(r"(20\d{2})")
RE_AUTH_FIN = re.compile(r"^\s*تفويض\s+مالي")
RE_AUTH_DEP = re.compile(r"^\s*تفويض\s+مصلحي")
RE_PERIOD = re.compile(r"من\s+([\d/ ]+?)\s+(?:الي|إلى|الى)\s+([\d/ ]+)")
RE_TRANSFER_OUT = re.compile(r"^\s*نقل\s+منه")
RE_TRANSFER_IN = re.compile(r"^\s*نقل\s+له")
RE_TARGET = re.compile(r"(\d+)\s*/\s*(\d+)|/\s*(\d+)")
RE_DEPOSIT = re.compile(r"الودائع|الامانات|الأمانات|صندوق")
SUMMARY_HINTS = ("اجمالى", "اجمالي", "إجمالي", "جملة")


def _dec(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, Decimal):
        return v
    if isinstance(v, float):
        return Decimal(repr(v)).quantize(Decimal("0.001"))
    try:
        return Decimal(str(v).replace(",", "").strip() or "0")
    except InvalidOperation:
        return Decimal("0")


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


@dataclass
class ParsedRow:
    sheet: str
    row_no: int
    date: str | None
    description: str
    doc: str | None
    amount: Decimal          # العمود D (زيادة/نقص الاعتماد)
    actual: Decimal          # العمود E
    commitment: Decimal      # العمود F
    notes: str | None
    excel_g: str | None      # قيمة Excel المحسوبة (للمقارنة فقط)
    excel_h: str | None
    classification: str = "UNCLASSIFIED"
    auth_type: str | None = None
    auth_no: str | None = None
    period_text: str | None = None
    target_item: str | None = None
    deposit_account: bool = False

    def as_json(self) -> dict:
        d = asdict(self)
        for k in ("amount", "actual", "commitment"):
            d[k] = str(d[k])
        return d


@dataclass
class ParsedSheet:
    name: str
    item_code: str | None
    year: int | None
    sector: str | None
    header_row: int | None
    rows: list[ParsedRow] = field(default_factory=list)
    constant_formula_cells: list[str] = field(default_factory=list)   # IMP-07
    last_excel_g: Decimal | None = None
    last_excel_h: Decimal | None = None


@dataclass
class SummaryLine:
    row_no: int
    item_code: str
    label: str
    formula: str | None
    value: Decimal | None


@dataclass
class ParsedSummary:
    name: str
    lines: list[SummaryLine]
    total_formula: str | None
    total_value: Decimal | None
    total_cell: str | None
    header_labels: list[str]


@dataclass
class ParsedWorkbook:
    sheets: list[ParsedSheet]
    summary: ParsedSummary | None


def classify(r: ParsedRow) -> None:
    b = (r.description or "").strip()
    if not b and r.amount == 0 and r.actual == 0 and r.commitment == 0:
        r.classification = "EMPTY_TEMPLATE_ROW"
        return
    if RE_AUTH_FIN.match(b) and r.amount > 0:
        r.classification, r.auth_type = "AUTHORIZATION_ALLOCATION", "FINANCIAL"
    elif RE_AUTH_DEP.match(b) and r.amount > 0:
        r.classification, r.auth_type = "AUTHORIZATION_ALLOCATION", "DEPARTMENTAL"
    elif RE_TRANSFER_OUT.match(b) and r.amount < 0:
        r.classification = "TRANSFER_OUT"
        m = RE_TARGET.search(b.split("بند", 1)[-1])
        if m:
            r.target_item = f"{m.group(1)}/{m.group(2)}" if m.group(1) else f"2/{m.group(3)}"
    elif RE_TRANSFER_IN.match(b) and r.amount > 0:
        r.classification = "TRANSFER_IN_AGGREGATE"
    elif r.commitment > 0:
        r.classification = "COMMITMENT"
    elif r.amount == 0 and r.actual > 0:
        r.classification = "ACTUAL_EXPENDITURE"
        r.deposit_account = bool(RE_DEPOSIT.search(b))
    if r.classification == "AUTHORIZATION_ALLOCATION":
        r.auth_no = r.doc
        m = RE_PERIOD.search(b)
        r.period_text = f"{m.group(1).strip()} — {m.group(2).strip()}" if m else None


def parse_period(text: str | None, year: int) -> tuple[date | None, date | None]:
    """«1/1 — 30/ 6» يوم/شهر، و«7/1 — 12/31» شهر/يوم. يُقبل التفسير الوحيد الصالح داخل السنة فقط."""
    if not text:
        return None, None
    parts = [p.strip() for p in text.split("—")]
    if len(parts) != 2:
        return None, None

    def nums(p):
        return [int(x) for x in re.findall(r"\d+", p)]

    a, b = nums(parts[0]), nums(parts[1])
    if len(a) < 2 or len(b) < 2 or (len(a) > 2 and a[2] != year) or (len(b) > 2 and b[2] != year):
        return None, None
    options = []
    for dm in (True, False):
        try:
            s = date(year, a[1], a[0]) if dm else date(year, a[0], a[1])
            e = date(year, b[1], b[0]) if dm else date(year, b[0], b[1])
            if s <= e:
                options.append((s, e))
        except ValueError:
            continue
    options = list(dict.fromkeys(options))
    return options[0] if len(options) == 1 else (None, None)


def _find_summary(ws, wsv) -> ParsedSummary | None:
    # ورقة الملخص: عنوانها يحتوي «إجمالي» ولا تحتوي رأس بند
    has_item_header = any(isinstance(c.value, str) and RE_ITEM.search(c.value)
                          for row in ws.iter_rows(max_row=20) for c in row)
    if has_item_header or not any(normalize(h) in normalize(ws.title) for h in SUMMARY_HINTS):
        return None
    lines, total_formula, total_value, total_cell, headers = [], None, None, None, []
    for row in ws.iter_rows():
        cells = {c.column_letter: c for c in row if hasattr(c, "column_letter")}
        texts = [str(c.value) for c in row if isinstance(c.value, str) and not str(c.value).startswith("=")]
        b, c_ = cells.get("B"), cells.get("C")
        e = cells.get("E")
        if b is not None and isinstance(b.value, (int, float)) and c_ is not None and isinstance(c_.value, (int, float)):
            d_cell = cells.get("D")
            lines.append(SummaryLine(b.row, f"{int(b.value)}/{int(c_.value)}",
                                     str(d_cell.value if d_cell is not None and d_cell.value else "").strip(),
                                     e.value if e is not None and isinstance(e.value, str) else None,
                                     _dec(wsv[e.coordinate].value) if e is not None else None))
        elif e is not None and isinstance(e.value, str) and e.value.upper().startswith("=SUM"):
            total_formula, total_cell = e.value, e.coordinate
            total_value = _dec(wsv[e.coordinate].value)
        elif texts and not lines:
            headers += texts
    return ParsedSummary(ws.title, lines, total_formula, total_value, total_cell, headers)


def parse_workbook(data: bytes) -> ParsedWorkbook:
    wf = load_workbook(io.BytesIO(data), data_only=False)
    wv = load_workbook(io.BytesIO(data), data_only=True)
    sheets, summary = [], None
    for ws in wf.worksheets:
        wsv = wv[ws.title]
        s = _find_summary(ws, wsv)
        if s is not None:
            summary = s
            continue
        item_code = year = sector = header_row = None
        data_start = None
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if item_code is None and (m := RE_ITEM.search(v)):
                    item_code = f"{m.group(1)}/{m.group(2)}"
                if "قطاع" in v or "مراقبة" in v:
                    sector = sector or (v if "مراقبة" in v else None)
                if (m := RE_YEAR.search(v)) and year is None and ("مراقبة" in v or "قطاع" in v):
                    year = int(m.group(1))
                if v.strip() == "رقم المستند":
                    header_row = c.row
                if v.strip() == "درهم دينار" and c.column_letter == "D":
                    data_start = c.row + 1
        if sector is None:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and "مراقبة" in c.value:
                        sector = c.value
                        m = RE_YEAR.search(c.value)
                        year = year or (int(m.group(1)) if m else None)
        ps = ParsedSheet(ws.title, item_code, year, sector, header_row)
        if data_start is None:
            sheets.append(ps)
            continue
        for r in range(data_start, ws.max_row + 1):
            def cell(col, _ws=ws, _r=r):
                return _ws[f"{col}{_r}"].value

            def vcell(col, _wsv=wsv, _r=r):
                return _wsv[f"{col}{_r}"].value

            desc = cell("B")
            doc = cell("C")
            row = ParsedRow(sheet=ws.title, row_no=r, date=(_date(cell("A")).isoformat() if _date(cell("A")) else None),
                            description=str(desc).strip() if desc is not None else "",
                            doc=str(doc).strip() if doc is not None else None, amount=_dec(cell("D")),
                            actual=_dec(cell("E")), commitment=_dec(cell("F")),
                            notes=str(cell("I")).strip() if cell("I") is not None else None,
                            excel_g=str(_dec(vcell("G"))), excel_h=str(_dec(vcell("H"))))
            classify(row)
            ps.rows.append(row)
            # IMP-07: قيمة ثابتة داخل عمود معادلات
            for col in ("G", "H"):
                v = cell(col)
                if v is not None and not (isinstance(v, str) and v.startswith("=")):
                    ps.constant_formula_cells.append(f"{col}{r}")
        meaningful = [x for x in ps.rows if x.classification != "EMPTY_TEMPLATE_ROW"] or ps.rows
        if meaningful:
            ps.last_excel_g, ps.last_excel_h = _dec(meaningful[-1].excel_g), _dec(meaningful[-1].excel_h)
        sheets.append(ps)
    return ParsedWorkbook(sheets, summary)
