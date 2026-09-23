"""محرك التقارير (المرحلة 11): التقارير العشرون، ومطابقة الأرقام للقيود، والتصدير."""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.db import new_session
from tests.ledger_helpers import actual, budget_doc, build_world, post, spec, transfer

API = "/api/v1"


@pytest.fixture()
def world(client):
    w = build_world(basis="APPROPRIATION", items=("2/6", "2/16", "2/18"))
    budget_doc(w, "ORIGINAL_BUDGET", {"2/18": "180000", "2/16": "40000", "2/6": "20000"}, on=date(2026, 1, 5))
    transfer(w, "2/6", "2/16", "20000")
    actual(w, "2/18", "150000")
    post(w, spec(w, "2/16", "COMMITMENT", "COMMITMENT", "50000"))
    post(w, spec(w, "2/18", "ACTUAL_EXPENDITURE", "ACTUAL", "1000"), on=date(2026, 3, 20), date_is_estimated=True)
    return w


def run(client, h, code, **params):
    r = client.post(f"{API}/reports/{code}/run", headers=h, json=params)
    assert r.status_code == 200, r.text
    return r.json()


def test_catalog_lists_20_reports_filtered_by_permission(client, user_factory):
    rv = user_factory("viewer", "REPORT_VIEWER")
    aud = user_factory("auditor", "AUDITOR")
    codes = [r["code"] for r in client.get(f"{API}/reports", headers=aud).json()]
    assert len(codes) == 20 and "RPT-16" in codes
    assert "RPT-16" not in [r["code"] for r in client.get(f"{API}/reports", headers=rv).json()]


def test_central_report_matches_ledger(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = run(client, h, "RPT-02", fiscal_year_id=str(world.fy_id))
    by = {x["item_code"]: x for x in r["rows"]}
    assert (by["2/16"]["budget"], by["2/16"]["transfers"], by["2/16"]["commitments"], by["2/16"]["available"]) == \
        ("40000.000", "20000.000", "50000.000", "10000.000")
    assert by["2/18"]["execution_rate"] == "83.89"
    with new_session() as s:
        actual_total = s.scalar(text("SELECT sum(direction*amount) FROM ledger_entries WHERE component='ACTUAL'"))
    assert Decimal(r["totals"]["actual"]) == actual_total   # Transactions → Calculations → Reports
    assert len(r["fingerprint"]) == 64


@pytest.mark.parametrize("code,extra", [
    ("RPT-01", {}), ("RPT-01", {"as_of": "2026-02-01"}), ("RPT-03", {}), ("RPT-04", {}), ("RPT-05", {}),
    ("RPT-06", {}), ("RPT-07", {}), ("RPT-08", {}), ("RPT-09", {}), ("RPT-10", {}), ("RPT-12", {}),
    ("RPT-13", {"month": 3}), ("RPT-14", {"quarter": 1}), ("RPT-15", {}), ("RPT-16", {}), ("RPT-17", {}),
    ("RPT-18", {}), ("RPT-19", {}), ("RPT-20", {"group_by": ["item", "txn_type"]}),
])
def test_every_report_runs(client, user_factory, world, code, extra):
    h = user_factory("aud", "AUDITOR")
    r = run(client, h, code, fiscal_year_id=str(world.fy_id), **extra)
    assert r["columns"] and r["title"]


def test_item_movement_and_as_of(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = run(client, h, "RPT-11", fiscal_year_id=str(world.fy_id), budget_line_id=str(world.lines["2/18"]))
    assert [x["available_after"] for x in r["rows"]] == ["180000.000", "30000.000", "29000.000"]
    early = run(client, h, "RPT-01", fiscal_year_id=str(world.fy_id), as_of="2026-01-31")
    assert {x["item_code"]: x["available"] for x in early["rows"]}["2/18"] == "180000.000"


def test_monthly_report_flags_estimated_dates(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = run(client, h, "RPT-13", fiscal_year_id=str(world.fy_id), month=3)
    assert any("تقديري" in n for n in r["notes"])
    assert {x["item_code"]: x["period_actual"] for x in r["rows"]}["2/18"] == "151000.000"


def test_custom_report_rejects_arbitrary_grouping(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = client.post(f"{API}/reports/RPT-20/run", headers=h,
                    json={"fiscal_year_id": str(world.fy_id), "group_by": ["le.amount; DROP TABLE users"]})
    assert r.status_code == 422 and r.json()["code"] == "INVALID_GROUP"


def test_audit_report_requires_audit_permission(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    assert client.post(f"{API}/reports/RPT-16/run", headers=h, json={}).status_code == 403


def test_pdf_export_with_embedded_arabic_font(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = client.post(f"{API}/reports/RPT-02/export", headers=h, params={"format": "pdf"},
                    json={"fiscal_year_id": str(world.fy_id)})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF")
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(r.content))
    fonts = {f.get_object()["/BaseFont"] for pg in reader.pages for f in pg["/Resources"]["/Font"].values()}
    assert any("Plex" in f for f in fonts)            # الخط العربي مضمَّن
    assert "RPT-02" in reader.pages[0].extract_text()
    with new_session() as s:
        row = s.execute(text("SELECT user_id, new_values FROM audit_log WHERE action='EXPORT'")).one()
    assert row.new_values["format"] == "pdf" and len(row.new_values["fingerprint"]) == 64


def test_xlsx_export_has_real_numbers_and_rtl(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = client.post(f"{API}/reports/RPT-02/export", headers=h, params={"format": "xlsx"},
                    json={"fiscal_year_id": str(world.fy_id)})
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["RPT-02"]
    assert ws.sheet_view.rightToLeft
    header = [c.value for c in ws[5]]
    assert header[:3] == ["البند", "اسم البند", "الاعتماد"]
    values = [row for row in ws.iter_rows(min_row=6, values_only=True) if row[0] == "2/16"][0]
    assert isinstance(values[2], (int, float)) and values[2] == 40000
    assert str(ws.cell(ws.max_row, 3).value).startswith("=SUM(")
    assert "المعاملات" in wb.sheetnames


def test_print_html(client, user_factory, world):
    h = user_factory("viewer", "REPORT_VIEWER")
    r = client.post(f"{API}/reports/RPT-07/export", headers=h, params={"format": "html"},
                    json={"fiscal_year_id": str(world.fy_id)})
    assert r.status_code == 200 and 'dir="rtl"' in r.text and "window.print" in r.text
