"""محرك الاستيراد على ملف اصطناعي مطابق للملف الحقيقي بنية وأرقامًا وأخطاءً (المرحلة 12)."""
from decimal import Decimal

import pytest
from sqlalchemy import select

from app.core.db import new_session
from app.modules.expenditures.models import Expenditure
from app.modules.imports.parser import parse_period
from tests.import_helpers import API, run_import
from tests.imports.synthetic import build

D = Decimal


@pytest.fixture(scope="module")
def workbook_bytes():
    return build()


def test_full_pipeline_reproduces_blueprint_analysis(client, team, user_factory, settings_tmp_storage, workbook_bytes):
    res = run_import(client, team, user_factory, workbook_bytes)
    # فصل المهام: المحضّر لا ينفذ، والتنفيذ يتطلب قبول الاستثناءات التاريخية أولًا
    assert res["first_commit"].json()["code"] == "HISTORICAL_EXCEPTIONS_NOT_ACCEPTED"
    assert res["own_commit"].status_code == 403
    counts = res["analysis"]["classification_counts"]
    counts.pop("EMPTY_TEMPLATE_ROW")
    assert counts == {"AUTHORIZATION_ALLOCATION": 20, "TRANSFER_OUT": 16, "TRANSFER_IN_AGGREGATE": 1,
                      "ACTUAL_EXPENDITURE": 24}
    codes = res["issue_codes"]
    for code in ("IMP-01", "IMP-05", "IMP-06", "IMP-07", "IMP-08", "IMP-09", "IMP-10", "IMP-11", "IMP-12", "IMP-13",
                 "IMP-14", "IMP-15", "IMP-16", "IMP-17", "DQ-04", "IMP-04", "IMP-NOTE", "IMP-DEP"):
        assert code in codes, code
    assert codes["IMP-07"] == 2 and codes["IMP-11"] == 2 and codes["IMP-15"] == 1
    msgs = {i["code"]: i["message"] for i in res["issues"]}
    assert "63,300.000" in msgs["IMP-05"]                          # 19,800 + 18,500 + 25,000
    assert "2/22" in [i["message"] for i in res["issues"] if i["code"] == "IMP-16"][0]

    assert res["created"] == {"authorizations": 5, "transfers": 16, "expenditures": 24}
    pos = res["positions"]
    assert (pos["2/16"]["transfer_in"], pos["2/16"]["available"]) == (D("80130"), D("-18370"))       # D-03
    assert pos["2/17"]["available"] == D("-1255")
    assert pos["2/18"]["available"] == D("-24969.420")
    assert pos["2/25"]["available"] == D("-13000")
    assert sum(p["actual"] for p in pos.values()) == D("498674.420")
    assert sum(p["allocation"] for p in pos.values()) == D("346034")
    assert sum(p["transfer_in"] for p in pos.values()) == sum(p["transfer_out"] for p in pos.values())


def test_preview_reconciliation_and_dq_report(client, team, user_factory, settings_tmp_storage, workbook_bytes):
    res = run_import(client, team, user_factory, workbook_bytes)
    rec = {r["item_code"]: r for r in res["preview"]["positions"]}
    assert (rec["2/18"]["excel_balance"], rec["2/18"]["book_balance"]) == ("30.580", "-24969.420")
    assert rec["2/16"]["difference"] == "-18370.000"
    r = client.get(f"{API}/imports/{res['batch_id']}/report", headers=res["prep"], params={"format": "pdf"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    j = client.get(f"{API}/imports/{res['batch_id']}/report", headers=res["prep"]).json()
    assert j["issues"][0]["severity"] == "حرج"


def test_imported_documents_are_flagged_and_traceable(client, team, user_factory, settings_tmp_storage, workbook_bytes):
    res = run_import(client, team, user_factory, workbook_bytes)
    with new_session() as s:
        exps = s.scalars(select(Expenditure)).all()
    assert all(e.status == "POSTED" and e.is_historical_exception and e.date_is_estimated for e in exps)
    dep = [e for e in exps if e.payment_method == "DEPOSIT_ACCOUNT"]
    assert len(dep) == 4                                             # D-07
    assert {e.document_no for e in exps if e.document_no.startswith("2/7-")} == {f"2/7-{i}" for i in range(1, 6)}
    assert all("!" in e.legacy_ref for e in exps)
    over = client.post(f"{API}/reports/RPT-08/run", headers=res["admin"],
                       json={"fiscal_year_id": str(res["world"].fy_id)}).json()
    assert {r["cause"] for r in over["rows"]} == {"استثناء تاريخي (مستورد)"}
    assert len(over["rows"]) == 9                                    # D-02: تسعة بنود متجاوزة


def test_same_file_cannot_be_imported_twice(client, team, user_factory, settings_tmp_storage, workbook_bytes):
    res = run_import(client, team, user_factory, workbook_bytes)
    r = client.post(f"{API}/imports", headers=res["prep"], files={"file": ("again.xlsx", workbook_bytes)},
                    data={"fiscal_year_id": str(res["world"].fy_id)})
    assert r.status_code == 409 and r.json()["code"] == "ALREADY_IMPORTED"


def test_authorization_amount_decision_and_row_dates(client, team, user_factory, settings_tmp_storage, workbook_bytes):
    run_import(client, team, user_factory, workbook_bytes, decisions={
        "authorization_amounts": {"FINANCIAL:479": "300000"},
        "row_dates": {"الصياتة !12": "2023-08-15"}})
    with new_session() as s:
        e = s.scalar(select(Expenditure).where(Expenditure.legacy_ref == "الصياتة !12"))
    assert str(e.expenditure_date) == "2023-08-15" and e.date_is_estimated is False


def test_bad_files_rejected(client, team, settings_tmp_storage):
    from tests.ledger_helpers import build_world
    w = build_world(year=2023, basis="AUTHORIZATION", items=())
    r = client.post(f"{API}/imports", headers=team["DATA_ENTRY"], files={"file": ("x.pdf", b"%PDF-1.4")},
                    data={"fiscal_year_id": str(w.fy_id)})
    assert r.status_code == 422
    r = client.post(f"{API}/imports", headers=team["DATA_ENTRY"], files={"file": ("x.xlsx", b"not a zip")},
                    data={"fiscal_year_id": str(w.fy_id)})
    assert r.status_code == 422


def test_unclassified_row_blocks_import(client, team, user_factory, settings_tmp_storage):
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build()))
    ws = wb["بريد"]
    ws["B12"], ws["D12"], ws["E12"] = "حركة غريبة", 5, 3
    buf = io.BytesIO()
    wb.save(buf)
    from tests.ledger_helpers import build_world
    w = build_world(year=2023, basis="AUTHORIZATION", items=())
    prep = team["DATA_ENTRY"]
    bid = client.post(f"{API}/imports", headers=prep, files={"file": ("x.xlsx", buf.getvalue())},
                      data={"fiscal_year_id": str(w.fy_id)}).json()["id"]
    client.post(f"{API}/imports/{bid}/analyze", headers=prep)
    stats = client.post(f"{API}/imports/{bid}/validate", headers=prep).json()
    assert stats["blocking"] is True
    client.put(f"{API}/imports/{bid}/decisions", headers=prep, json={"accept_historical_exceptions": True})
    client.post(f"{API}/imports/{bid}/validate", headers=prep)
    r = client.post(f"{API}/imports/{bid}/commit", headers=user_factory("adm2", "SYSTEM_ADMIN"))
    assert r.status_code == 409 and r.json()["code"] == "BLOCKING_ISSUES"


def test_period_parsing_rules():
    from datetime import date
    assert parse_period("1/1 — 30/ 6", 2023) == (date(2023, 1, 1), date(2023, 6, 30))       # يوم/شهر
    assert parse_period("7/1 — 12/31", 2023) == (date(2023, 7, 1), date(2023, 12, 31))      # شهر/يوم
    assert parse_period("7/1 — 31/ 12 /2020", 2023) == (None, None)
